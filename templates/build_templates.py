#!/usr/bin/env python3
"""주식일지 Google Sheets 양식(.xlsx) 생성기.

실제 거래·금액·시트 ID 없이, 운용 중인 3개 파일의 구조·수식·검증 규칙을 그대로 옮기고
가상의 샘플 데이터만 넣습니다. 생성된 .xlsx를 Google Drive에 올려 Google Sheets로 열면
(파일 → 저장: Google Sheets로) 그대로 사용할 수 있습니다.

    pip install openpyxl
    python3 templates/build_templates.py        # templates/*.xlsx 다시 생성

시트 구조를 바꿀 때는 이 스크립트도 함께 고쳐 주세요.
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent

# ── 서식 ──────────────────────────────────────────────
KRW = '₩#,##0;-₩#,##0'
USD = '$#,##0.00;-$#,##0.00'
PCT = '0.00%'
QTY = '#,##0.####'
DATE = 'yyyy-mm-dd'
FX = '#,##0.00'

TITLE = Font(bold=True, size=14)
NOTE = Font(italic=True, size=9, color='666666')
BOLD = Font(bold=True)
INPUT = Font(color='1155CC')  # 파란 글씨 = 직접 입력
HEAD_FILL = PatternFill('solid', fgColor='E8EEF7')
TOTAL_FILL = PatternFill('solid', fgColor='F3F3F3')
THIN = Side(style='thin', color='BBBBBB')
BOX = Border(bottom=THIN)

KR_ACCOUNTS = ['메리츠증권', '메리츠 RIA', '토스증권', '토스 RIA']
US_ACCOUNTS = ['메리츠증권', '토스증권']
FLOW_TYPES = ['입금', '출금', '달러 매수', '달러 매도']


def header(ws, row, labels, col=1):
    for i, v in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def fmt(ws, rng, number_format):
    for row in ws[rng]:
        for c in row:
            c.number_format = number_format


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def list_dv(ws, items, rng, strict=True):
    dv = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True,
                        showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(rng)


# ══════════════════════════════════════════════════════
# 1. 주식일지 (메인)
# ══════════════════════════════════════════════════════
def build_main():
    wb = Workbook()
    wb.remove(wb.active)

    # ── 총자산 ──
    ws = wb.create_sheet('총자산')
    ws.sheet_view.showGridLines = False
    ws['A1'] = '주식 총자산'; ws['A1'].font = TITLE
    ws['A2'] = ("종목별 보유 현황은 '국장투자' / '미장투자' 시트, 투자 원금·예수금은 '현금투자' 시트에서 확인하세요. "
                "환율·파일 URL은 '설정' 시트에서 변경합니다.")
    ws['A2'].font = NOTE
    header(ws, 3, ['요약', '국장 (KRW)', '미장 (USD)', '미장 (원화 환산)', '합계 (KRW)', '합계 (USD 환산)'])
    rows = [
        ('투자 원금', "='현금투자'!B16", "='현금투자'!B29", "='현금투자'!B31"),
        ('매수금액', "='국장투자'!G4", "='미장투자'!G4", None),
        ('평가금액', "='국장투자'!H4", "='미장투자'!H4", None),
        ('예수금', "='현금투자'!C13", "='현금투자'!C26", None),
        ('총 평가자산', '=B6+B7', '=C6+C7', '=D6+D7'),
        ('총 손익', '=B8-B4', '=C8-C4', '=D8-D4'),
        ('수익률', '=IF(B4=0,0,B9/B4)', '=IF(C4=0,0,C9/C4)', '=IF(D4=0,0,D9/D4)'),
        ('누적 실현손익', "='국장투자'!K4", "='미장투자'!K4", None),
    ]
    for i, (label, b, c, d) in enumerate(rows):
        r = 4 + i
        ws.cell(r, 1, label).font = BOLD
        ws.cell(r, 2, b)
        ws.cell(r, 3, c)
        ws.cell(r, 4, d or f'=C{r}*FX_USDKRW')
        if label == '수익률':
            ws.cell(r, 5, f'=IF(E4=0,0,E9/E4)')
            ws.cell(r, 6, '=E10')
        else:
            ws.cell(r, 5, f'=B{r}+D{r}' if label not in ('총 평가자산', '총 손익') else
                    ('=E6+E7' if label == '총 평가자산' else '=E8-E4'))
            ws.cell(r, 6, f'=E{r}/FX_USDKRW')
        for col, nf in zip('BCDEF', [KRW, USD, KRW, KRW, USD]):
            ws[f'{col}{r}'].number_format = PCT if label == '수익률' else nf
    ws['A13'] = ("투자 원금은 '현금투자' 탭의 입출금·환전 내역 누적, 예수금은 기준일 스냅샷 + 이후 매매·입출금·환전 기준입니다. "
                 "누적 실현손익 행은 참고용 전체 기간 합계입니다. 미장 원화 환산 투자 원금은 누적 달러 매수 평균 환율"
                 "(현금투자 B30) 기준이라 총 손익에 환차손익이 포함됩니다.")
    ws['A13'].font = NOTE
    ws.freeze_panes = 'B1'
    widths(ws, {'A': 16, 'B': 16, 'C': 14, 'D': 16, 'E': 16, 'F': 14})

    # ── 국장투자 / 미장투자 ──
    for title, label, url, money in (('국장투자', '국장 보유 종목 (KRW)', 'URL_KR', KRW),
                                      ('미장투자', '미장 보유 종목 (USD)', 'URL_US', USD)):
        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws['A1'] = label; ws['A1'].font = TITLE
        src = '국장_주식일지' if url == 'URL_KR' else '미장_주식일지'
        ws['A2'] = f"보유 종목은 {src}의 '종목목록' 탭에서 개수 제한 없이 자동으로 불러옵니다."
        ws['A2'].font = NOTE
        header(ws, 3, ['종목 탭', '종목명', '티커', '보유 수량', '평균 매수단가', '현재가', '매수금액',
                       '평가금액', '평가 손익', '수익률', '누적 실현손익', '비중'])
        ws['A4'] = '합계'
        ws['B4'] = '=COUNTIF(A5:A,"?*")&"종목"'
        ws['G4'] = '=SUM(G5:G)'
        ws['H4'] = '=SUM(H5:H)'
        ws['I4'] = '=H4-G4'
        ws['J4'] = '=IF(G4=0,0,I4/G4)'
        ws['K4'] = '=SUM(K5:K)'
        ws['L4'] = '=SUM(L5:L)'
        for c in ws[4]:
            c.font = BOLD; c.fill = TOTAL_FILL
        ws['A5'] = (
            f'=ARRAYFORMULA(IFERROR(LET(tl,IMPORTRANGE({url},"\'종목목록\'!A2:A"),tabs,FILTER(tl,tl<>""),'
            f'rs,REDUCE(SEQUENCE(1,11,0,0),tabs,LAMBDA(acc,t,VSTACK(acc,LET(d,IMPORTRANGE({url},"\'"&t&"\'!A1:D7"),'
            'HSTACK(t,INDEX(d,1,2),INDEX(d,2,2),INDEX(d,4,2),INDEX(d,5,2),INDEX(d,3,2),INDEX(d,6,2),INDEX(d,7,2),'
            'INDEX(d,1,4),INDEX(d,2,4),INDEX(d,3,4)))))),r,FILTER(rs,SEQUENCE(ROWS(rs))>1),h,CHOOSECOLS(r,8),'
            'tot,SUM(h),HSTACK(r,IF(tot=0,"",h/tot))),""))')
        fmt(ws, 'D4:D300', QTY)
        for col in 'EFGHIK':
            fmt(ws, f'{col}4:{col}300', money)
        fmt(ws, 'J4:J300', PCT)
        fmt(ws, 'L4:L300', PCT)
        ws.freeze_panes = 'A5'
        widths(ws, {'A': 18, 'B': 18, 'C': 14, 'D': 10, 'E': 14, 'F': 13, 'G': 15, 'H': 15, 'I': 14,
                    'J': 9, 'K': 14, 'L': 8})

    # ── 현금투자 ──
    ws = wb.create_sheet('현금투자')
    ws.sheet_view.showGridLines = False
    ws['A1'] = '현금투자 현황'; ws['A1'].font = TITLE
    ws['A2'] = ("투자 원금 = 하단 '입출금 · 환전 내역' 누적(자동). 원화는 입금·출금으로, 달러는 환전(달러 매수/매도)으로만 "
                "들어오고 나갑니다. 미장 원화 원금 = 달러 원금 × 기준 환율(누적 달러 매수 평균, 자동). 예수금 = 기준 예수금 + "
                "기준일 이후 매도대금 − 매수대금 + 입출금·환전(자동). 기준일 당일 행은 원금 이월분이라 예수금에는 더해지지 "
                "않습니다. 파란 칸만 직접 입력합니다.")
    ws['A2'].font = NOTE
    ws['A3'] = '기준일'; ws['A3'].font = BOLD
    ws['B3'] = date(2026, 1, 2); ws['B3'].number_format = DATE; ws['B3'].font = INPUT

    agg = "'_거래집계'!"
    def acct_rows(first, mkt, accounts, base_cash, total_row, money, usd):
        for i in range(6):
            r = first + i
            a = f'$A{r}'
            if i < len(accounts):
                ws.cell(r, 1, accounts[i]).font = INPUT
                ws.cell(r, 6, base_cash[i]).font = INPUT
            ws.cell(r, 2, f'=IF({a}="","",SUMIFS({agg}$D:$D,{agg}$A:$A,"{mkt}",{agg}$F:$F,{a},{agg}$C:$C,"매수")'
                          f'-SUMIFS({agg}$D:$D,{agg}$A:$A,"{mkt}",{agg}$F:$F,{a},{agg}$C:$C,"매도")'
                          f'+SUMIFS({agg}$E:$E,{agg}$A:$A,"{mkt}",{agg}$F:$F,{a}))')
            ws.cell(r, 3, f'=IF({a}="","",N($F{r})+SUMIFS({agg}$D:$D,{agg}$A:$A,"{mkt}",{agg}$F:$F,{a},'
                          f'{agg}$C:$C,"매도",{agg}$G:$G,">"&$B$3)-SUMIFS({agg}$D:$D,{agg}$A:$A,"{mkt}",'
                          f'{agg}$F:$F,{a},{agg}$C:$C,"매수",{agg}$G:$G,">"&$B$3)+N($H{r}))')
            ws.cell(r, 4, f'=IF({a}="","",B{r}+C{r})')
            ws.cell(r, 5, f'=IF(OR({a}="",$D${total_row}=0),"",D{r}/$D${total_row})')
            ws.cell(r, 7, f'=IF({a}="","",SUMIFS({agg}$E:$E,{agg}$A:$A,"{mkt}",{agg}$F:$F,{a},{agg}$G:$G,">"&$B$3))')
            if usd:
                ws.cell(r, 8, f'=IF({a}="","",SUMIFS($F$36:$F,$C$36:$C,{a},$B$36:$B,"달러 매수",$A$36:$A,">"&$B$3)'
                              f'-SUMIFS($F$36:$F,$C$36:$C,{a},$B$36:$B,"달러 매도",$A$36:$A,">"&$B$3))')
            else:
                ws.cell(r, 8, f'=IF({a}="","",SUMIFS($D$36:$D,$C$36:$C,{a},$B$36:$B,"입금",$A$36:$A,">"&$B$3)'
                              f'-SUMIFS($D$36:$D,$C$36:$C,{a},$B$36:$B,"출금",$A$36:$A,">"&$B$3)'
                              f'-SUMIFS($D$36:$D,$C$36:$C,{a},$B$36:$B,"달러 매수",$A$36:$A,">"&$B$3)'
                              f'+SUMIFS($D$36:$D,$C$36:$C,{a},$B$36:$B,"달러 매도",$A$36:$A,">"&$B$3))')
            for col in 'BCDFGH':
                ws[f'{col}{r}'].number_format = money
            ws[f'E{r}'].number_format = PCT
        last = first + 5
        ws.cell(total_row, 1, '합계')
        for col in 'BCDEFGH':
            ws[f'{col}{total_row}'] = f'=SUM({col}{first}:{col}{last})'
            ws[f'{col}{total_row}'].number_format = PCT if col == 'E' else money
        for c in ws[total_row][:8]:
            c.font = BOLD; c.fill = TOTAL_FILL

    ws['A5'] = '원화 계좌 (국장 · KRW)'; ws['A5'].font = Font(bold=True, size=12)
    header(ws, 6, ['계좌', '투자금 (매수원금)', '예수금', '계좌 합계', '비중', '기준 예수금',
                   '기준일 이후 실현손익', '기준일 이후 입출금·환전'])
    acct_rows(7, '국장', KR_ACCOUNTS, [7300000, 0, 3000000, 0], 13, KRW, usd=False)
    ws['A14'] = '누적 입금 − 출금 (원화)'
    ws['B14'] = '=SUMIFS($D$36:$D,$B$36:$B,"입금")-SUMIFS($D$36:$D,$B$36:$B,"출금")'
    ws['A15'] = '환전 순이동 (달러 매도분 × 기준 환율 − 달러 매수 원화)'
    ws['B15'] = '=-SUMIFS($D$36:$D,$B$36:$B,"달러 매수")+SUMIFS($F$36:$F,$B$36:$B,"달러 매도")*$B$30'
    ws['A16'] = '투자 원금'; ws['A16'].font = BOLD
    ws['B16'] = '=B14+B15'; ws['B16'].font = BOLD
    for r in (14, 15, 16):
        ws[f'B{r}'].number_format = KRW

    ws['A18'] = '달러 계좌 (미장 · USD)'; ws['A18'].font = Font(bold=True, size=12)
    header(ws, 19, ['계좌', '투자금 (매수원금)', '예수금', '계좌 합계', '비중', '기준 예수금',
                    '기준일 이후 실현손익', '기준일 이후 환전 (달러)'])
    acct_rows(20, '미장', US_ACCOUNTS, [2000, 0], 26, USD, usd=True)
    ws['A27'] = '누적 달러 매수'
    ws['B27'] = '=SUMIFS($F$36:$F,$B$36:$B,"달러 매수")'
    ws['A28'] = '누적 달러 매도'
    ws['B28'] = '=-SUMIFS($F$36:$F,$B$36:$B,"달러 매도")'
    ws['A29'] = '투자 원금'; ws['A29'].font = BOLD
    ws['B29'] = '=B27+B28'; ws['B29'].font = BOLD
    ws['A30'] = '기준 환율 (누적 달러 매수 평균 · 자동)'
    ws['B30'] = '=IF(B27=0,FX_USDKRW,SUMIFS($D$36:$D,$B$36:$B,"달러 매수")/B27)'
    ws['C30'] = '현재 환율 (설정 B5)'
    ws['D30'] = '=FX_USDKRW'
    ws['A31'] = '투자 원금 (원화 환산)'
    ws['B31'] = '=B29*B30'
    ws['C31'] = '환차손익 (미실현, 원금 기준)'
    ws['D31'] = '=B29*FX_USDKRW-B31'
    ws['A32'] = '목록에 없는 거래 계좌'
    ws['B32'] = ('=ARRAYFORMULA(LET(m,\'_거래집계\'!A2:A,f,\'_거래집계\'!F2:F,acc,UNIQUE(FILTER(m&" · "&f,m<>"")),'
                 'reg,VSTACK("국장 · "&A7:A12,"미장 · "&A20:A25),miss,FILTER(acc,ISNA(MATCH(acc,reg,0))),'
                 'IFERROR(TEXTJOIN(", ",TRUE,miss),"없음 ✓")))')
    ws['C32'] = '환전 실현손익 (달러 매도)'
    ws['D32'] = '=SUMIFS($D$36:$D,$B$36:$B,"달러 매도")-SUMIFS($F$36:$F,$B$36:$B,"달러 매도")*$B$30'
    for ref, nf in (('B27', USD), ('B28', USD), ('B29', USD), ('B30', FX), ('D30', FX),
                    ('B31', KRW), ('D31', KRW), ('D32', KRW)):
        ws[ref].number_format = nf

    ws['A34'] = '입출금 · 환전 내역 (원금 산정 기준)'; ws['A34'].font = Font(bold=True, size=12)
    header(ws, 35, ['일자', '구분', '계좌', '원화 금액', '적용 환율', '달러 금액', '메모'])
    flows = [  # 일자, 구분, 계좌, 원화, 환율, 달러, 메모 (샘플)
        (date(2026, 1, 2), '입금', '메리츠증권', 10000000, None, None, '샘플: 기준일 원금 이월'),
        (date(2026, 1, 2), '입금', '토스증권', 3000000, None, None, '샘플: 기준일 원금 이월'),
        (date(2026, 1, 2), '달러 매수', '메리츠증권', 2700000, 1350, 2000, '샘플: 기준일 달러 원금 이월'),
        (date(2026, 2, 2), '입금', '토스증권', 1000000, None, None, '샘플: 추가 입금'),
        (date(2026, 3, 3), '달러 매수', '토스증권', 676000, 1352, 500, '샘플: 환전'),
    ]
    for i, row in enumerate(flows):
        r = 36 + i
        for j, v in enumerate(row):
            if v is not None:
                c = ws.cell(r, 1 + j, v)
                c.font = INPUT
        ws[f'A{r}'].number_format = DATE
    fmt(ws, 'A36:A1000', DATE)
    fmt(ws, 'D36:D1000', KRW)
    fmt(ws, 'E36:E1000', FX)
    fmt(ws, 'F36:F1000', USD)
    list_dv(ws, FLOW_TYPES, 'B36:B1000')
    dv = DataValidation(type='list', formula1="='현금투자'!$A$7:$A$12", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv); dv.add('C36:C1000')
    widths(ws, {'A': 30, 'B': 16, 'C': 14, 'D': 16, 'E': 10, 'F': 14, 'G': 20, 'H': 20})

    # ── 설정 ──
    ws = wb.create_sheet('설정')
    ws.sheet_view.showGridLines = False
    ws['A1'] = '설정'; ws['A1'].font = TITLE
    ws['A2'] = '파란 글씨는 직접 입력하는 칸입니다. 모든 시트는 이름 범위(FX_USDKRW, URL_KR, URL_US)로 이 값을 참조합니다.'
    ws['A2'].font = NOTE
    header(ws, 4, ['항목', '값', '이름 범위', '설명'])
    ws.append(['환율 (USD/KRW)', '=IFERROR(GOOGLEFINANCE("CURRENCY:USDKRW"), 1350)', 'FX_USDKRW',
               'GOOGLEFINANCE 실시간 환율 (실패 시 기본값 1,350)'])
    ws.append(['국장 파일 URL', 'https://docs.google.com/spreadsheets/d/<국장_주식일지 ID>', 'URL_KR',
               '국장_주식일지 스프레드시트 주소'])
    ws.append(['미장 파일 URL', 'https://docs.google.com/spreadsheets/d/<미장_주식일지 ID>', 'URL_US',
               '미장_주식일지 스프레드시트 주소'])
    ws['B5'].number_format = FX
    ws['B6'].font = INPUT; ws['B7'].font = INPUT
    widths(ws, {'A': 18, 'B': 62, 'C': 14, 'D': 44})
    for name, ref in (('FX_USDKRW', '$B$5'), ('URL_KR', '$B$6'), ('URL_US', '$B$7')):
        wb.defined_names[name] = DefinedName(name, attr_text=f"'설정'!{ref}")

    # ── _거래집계 (숨김) ──
    ws = wb.create_sheet('_거래집계')
    ws['A1'] = (
        '=ARRAYFORMULA(LET(imp,LAMBDA(url,mkt,LET(tl,IMPORTRANGE(url,"\'종목목록\'!A2:A"),tabs,FILTER(tl,tl<>""),'
        'REDUCE({"","","","","","",""},tabs,LAMBDA(acc,t,IFERROR(LET(d,IMPORTRANGE(url,"\'"&t&"\'!A9:J"),'
        'f,FILTER(d,(INDEX(d,,2)="매수")+(INDEX(d,,2)="매도")),VSTACK(acc,HSTACK(IF(INDEX(f,,2)<>"",mkt),'
        'IF(INDEX(f,,2)<>"",t),CHOOSECOLS(f,2,6,9,10,1)))),acc))))),all,VSTACK(imp(URL_KR,"국장"),imp(URL_US,"미장")),'
        'VSTACK({"시장","종목","구분","정산금액","실현손익","계좌","일자"},FILTER(all,INDEX(all,,1)<>""))))')
    ws.sheet_state = 'hidden'

    # ── App_* (웹앱 전용, 헤더만) ──
    app = {
        'App_매매계획': ['종목코드', '시장', '종목탭', '목표가', '손절가', '목표비중(%)', '투자근거', '메모', '수정일'],
        'App_매매회고': ['거래키', '시장', '종목탭', '거래일자', '구분', '수량', '단가', '전략태그', '감정', '회고', '평점', '작성일'],
        'App_자산추이': ['일자', '투자원금(KRW)', '총평가자산(KRW)', '총손익(KRW)', '수익률', '국장 총평가(KRW)',
                     '미장 총평가(USD)', '예수금(KRW환산)', '환율(USD/KRW)', '국장 원금(KRW)', '미장 원금(USD)', '기록시각'],
        'App_관심종목': ['티커', '종목명', '시장', '메모', '추가일'],
        'App_전망': ['작성일', '구분', '티커', '종목명', '단기', '중기', '장기', '요약', '근거·이벤트', '출처'],
        'App_일간전망': ['작성일', '구분', '티커', '종목명', '오늘 전망', '요약', '체크포인트', '출처'],
        'App_시그널': ['키', '신호일', '시장', '티커', '종목명', '구분', '전략', '신호', '근거', '가격', '기준 심볼', '기록시각'],
        'App_재무추이': ['키', '일자', '시장', '티커', '종목명', '현재가', '실적 PER', '12개월 선행 PER', '포워드 PER',
                     '실적 EPS', '12개월 선행 EPS', '올해 회계연도 말', '올해 추정 EPS', '내년 회계연도 말',
                     '내년 추정 EPS', '목표주가', '투자의견(5=강력매수)', '출처', '기록시각'],
    }
    for title, cols in app.items():
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font = BOLD
        ws.freeze_panes = 'A2'
    wb['App_관심종목'].append(['NASDAQ:MSFT', 'Microsoft', '미장', '샘플 관심종목', '2026-01-02'])

    wb.save(OUT / '주식일지_양식.xlsx')


# ══════════════════════════════════════════════════════
# 2·3. 국장_주식일지 / 미장_주식일지 (종목별 상세)
# ══════════════════════════════════════════════════════
def stock_tab(wb, title, name, ticker, fallback, trades, money, accounts):
    ws = wb.create_sheet(title)
    left = [('종목명', name), ('티커', ticker), ('현재가', f'=IFERROR(GOOGLEFINANCE(B2, "price"), {fallback})'),
            ('보유 수량', '=SUMIFS(C9:C,B9:B,"매수")-SUMPRODUCT((B9:B="매도")*ABS(C9:C))'),
            ('평균 매수단가', '=IFERROR(INDEX(H:H, MAX(FILTER(ROW(H9:H), H9:H<>""))), 0)'),
            ('총 매수금액', '=B4*B5'), ('현재 평가금액', '=B4*B3')]
    right = [('평가 손익', '=B7-B6'), ('수익률 %', '=IF(B6=0, 0, D1/B6)'), ('누적 실현손익', '=SUM(I9:I)'),
             ('총 누적손익', '=D1+D3')]
    for i, (k, v) in enumerate(left, 1):
        ws.cell(i, 1, k).font = BOLD
        ws.cell(i, 2, v)
    for i, (k, v) in enumerate(right, 1):
        ws.cell(i, 3, k).font = BOLD
        ws.cell(i, 4, v)
    ws['B1'].font = INPUT; ws['B2'].font = INPUT
    for ref in ('B3', 'B5', 'B6', 'B7', 'D1', 'D3', 'D4'):
        ws[ref].number_format = money
    ws['B4'].number_format = QTY
    ws['D2'].number_format = PCT

    header(ws, 8, ['거래일자', '구분', '수량', '거래단가', '수수료/세금', '정산금액', '보유수량 변화',
                   '평단가 변화', '실현손익', '계좌구분', '메모'])
    for i, (d, kind, qty, px, fee, acct, memo) in enumerate(trades):
        r = 9 + i
        for col, v in zip('ABCDE', (d, kind, qty, px, fee)):
            ws[f'{col}{r}'] = v
            ws[f'{col}{r}'].font = INPUT
        ws[f'F{r}'] = f'=IF($B{r}="","",IF($B{r}="매도",ABS($C{r})*$D{r}-$E{r},$C{r}*$D{r}+$E{r}))'
        ws[f'G{r}'] = f'=IF($B{r}="","",N(G{r-1})+IF($B{r}="매도",-ABS($C{r}),ABS($C{r})))'
        ws[f'H{r}'] = f'=IF($B{r}="","",IF($B{r}="매도",IF(G{r}=0,0,N(H{r-1})),(N(G{r-1})*N(H{r-1})+$F{r})/G{r}))'
        ws[f'I{r}'] = f'=IF($B{r}="","",IF($B{r}="매도",$F{r}-ABS($C{r})*N(H{r-1}),0))'
        ws[f'J{r}'] = acct; ws[f'J{r}'].font = INPUT
        if memo:
            ws[f'K{r}'] = memo; ws[f'K{r}'].font = INPUT
    fmt(ws, 'A9:A1000', DATE)
    fmt(ws, 'C9:C1000', QTY)
    fmt(ws, 'G9:G1000', QTY)
    for col in 'DEFHI':
        fmt(ws, f'{col}9:{col}1000', money)
    dv = DataValidation(type='date', operator='greaterThan', formula1='1', allow_blank=True,
                        promptTitle='거래일자', prompt='셀을 더블클릭하면 달력이 열립니다. 오늘 날짜: Ctrl+;')
    ws.add_data_validation(dv); dv.add('A9:A1000')
    list_dv(ws, ['매수', '매도'], 'B9:B1000')
    list_dv(ws, accounts, 'J9:J1000')
    widths(ws, {'A': 14, 'B': 16, 'C': 14, 'D': 14, 'E': 12, 'F': 14, 'G': 12, 'H': 14, 'I': 13, 'J': 13, 'K': 20})


def stock_book(filename, tabs, money, accounts):
    wb = Workbook()
    wb.remove(wb.active)
    for t in tabs:
        stock_tab(wb, *t, money=money, accounts=accounts)
    ws = wb.create_sheet('종목목록')
    ws['A1'] = '종목 탭'; ws['A1'].font = BOLD
    ws['B1'] = '※ 웹앱으로 새 종목 탭을 만들면 자동 추가 (시트에서 직접 탭을 만들었다면 여기에 탭 이름 추가)'
    ws['B1'].font = NOTE
    for t in tabs:
        ws.append([t[0]])
    ws.column_dimensions['A'].width = 20
    wb.save(OUT / filename)


def build_kr():
    stock_book('국장_주식일지_양식.xlsx', [
        ('NAVER', 'NAVER', 'KRX:035420', 210000, [
            (date(2026, 1, 12), '매수', 10, 200000, 0, '메리츠증권', '샘플'),
            (date(2026, 2, 10), '매수', 5, 185000, 0, '메리츠증권', '샘플: 물타기'),
            (date(2026, 3, 16), '매도', 4, 230000, 0, '메리츠증권', '샘플: 일부 매도'),
        ]),
        ('KODEX200', 'KODEX 200', 'KRX:069500', 38000, [
            (date(2026, 2, 3), '매수', 40, 36000, 216, '토스증권', '샘플: 수수료 0.015%'),
        ]),
    ], KRW, KR_ACCOUNTS)


def build_us():
    stock_book('미장_주식일지_양식.xlsx', [
        ('AAPL', 'Apple', 'NASDAQ:AAPL', 245, [
            (date(2026, 1, 20), '매수', 5, 230, 0, '메리츠증권', '샘플'),
            (date(2026, 3, 5), '매수', 2, 240, 0.48, '토스증권', '샘플: 수수료 0.1%'),
            (date(2026, 3, 20), '매도', 2, 250, 0, '메리츠증권', '샘플: 일부 매도'),
        ]),
    ], USD, US_ACCOUNTS)


if __name__ == '__main__':
    build_main()
    build_kr()
    build_us()
    print('생성 완료:', ', '.join(p.name for p in sorted(OUT.glob('*.xlsx'))))
